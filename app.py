import atexit
import base64
import copy
from flask import Flask, request, jsonify, render_template, send_file, redirect, url_for
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.pagebreak import Break
from openpyxl.worksheet.properties import PageSetupProperties
import io
import math
import re
import shutil
import unicodedata
from datetime import datetime
import threading
import time
from concurrent.futures import ThreadPoolExecutor
import google.generativeai as genai
import os
import requests
from dotenv import load_dotenv
import json
import uuid

app = Flask(__name__)
app.config["UPLOAD_FOLDER"] = "uploads/"
os.makedirs(app.config["UPLOAD_FOLDER"], exist_ok=True)
load_dotenv()

genai.configure(api_key=os.environ["GEMINI_API_KEY"])

system_instruction = "You are a highly accurate and adaptable data extraction expert specializing in Arabic invoice images. Given an invoice image, meticulously extract and return all relevant data in JSON format. Ensure the output is comprehensive, error-free, and formatted correctly, regardless of the image's layout or format. Adapt to different styles and structures to provide the most accurate and complete data extraction. If the image quality is poor, apply image enhancement techniques to improve clarity."

ALLOWED_MIME_TYPES = {
    "application/pdf", "text/plain", "text/html", "text/csv", "text/xml", "text/rtf",
    "image/jpeg", "image/png", "image/gif", "image/webp", "image/heic", "image/heif",
}

# A model is a list of fields the user defined (name, type, where to find it). The prompt sent to Gemini is generated from it.
FIELD_TYPES = {
    "text": "Chữ",
    "number": "Số",
    "money": "Tiền (VND)",
    "date": "Ngày",
    "percent": "Phần trăm (%)",
    "yesno": "Có / Không",
}
TYPE_INSTRUCTIONS = {
    "text": "a JSON string, exactly as printed",
    "number": "a JSON number without thousand separators, '.' as decimal point",
    "money": "a JSON number in VND without thousand separators, currency symbol or unit",
    "date": "a JSON string in DD/MM/YYYY format",
    "percent": "a JSON number, for example 10 for 10%",
    "yesno": "true or false",
}

# Where a field is printed on the form. The Excel export uses it to rebuild the layout of the original document.
ZONES = ("top_left", "top_right", "subtitle", "info", "totals", "footer", "bottom")

# Prefilled in the "create model" dialog so a new model does not start empty
DEFAULT_TITLE = "HÓA ĐƠN"
DEFAULT_SIGNATURES = ["NGƯỜI MUA HÀNG", "NGƯỜI BÁN HÀNG"]
DEFAULT_FIELDS = [
    {"label": "Tên người bán", "type": "text", "hint": "", "zone": "top_left", "required": True},
    {"label": "Mã số thuế người bán", "type": "text", "hint": "", "zone": "top_left", "required": False},
    {"label": "Số hóa đơn", "type": "text", "hint": "", "zone": "top_right", "required": True},
    {"label": "Ngày lập", "type": "date", "hint": "", "zone": "subtitle", "required": True},
    {"label": "Tên người mua", "type": "text", "hint": "", "zone": "info", "required": False},
    {"label": "Mã số thuế người mua", "type": "text", "hint": "", "zone": "info", "required": False},
    {"label": "Tổng tiền thanh toán", "type": "money", "hint": "", "zone": "totals", "required": True},
]
DEFAULT_LINE_FIELDS = [
    {"label": "STT", "type": "number", "hint": "", "required": False},
    {"label": "Tên hàng hóa, dịch vụ", "type": "text", "hint": "", "required": True},
    {"label": "Đơn vị tính", "type": "text", "hint": "", "required": False},
    {"label": "Số lượng", "type": "number", "hint": "", "required": False},
    {"label": "Đơn giá", "type": "money", "hint": "", "required": False},
    {"label": "Thành tiền", "type": "money", "hint": "", "required": True},
]

RULES_TEXT = """General rules:
- Numbers must be JSON numbers. Vietnamese format uses "." for thousands and "," for decimals: "164.617.992" means 164617992 and "119,23" means 119.23. Convert other conventions the same way.
- If a value is not on the document, use null. Never invent values.
- Read Vietnamese diacritics exactly. If a word is split across two lines, join it back.
- If quantity, unit price and amount are all printed, check that they agree and use that to resolve an unclear digit, but still report what is printed.
- Also add one extra key "_image_quality" to the JSON object: "clear" if the document is easy to read, "blurry" if some text is hard to read because the image is blurry, out of focus or very low resolution, or "unreadable" if most of the text cannot be read."""

SUGGEST_PROMPT = """Look at this sample invoice (or receipt, or delivery note). The user will process many documents of this same form, so list the data fields worth extracting from it and describe where each one is printed.

Return one JSON object: {"title": "...", "signatures": [...], "fields": [...], "line_fields": [...], "notes": "..."}
- title: the main title printed on the form (for example "HÓA ĐƠN GIÁ TRỊ GIA TĂNG"), exactly as printed.
- signatures: the signature captions printed at the bottom, in order, for example ["NGƯỜI MUA HÀNG", "NGƯỜI BÁN HÀNG"]. Use [] if there are none.
- fields: values printed once on the document, such as invoice number, date, seller, buyer, tax code, totals, payment method, vehicle plate, cashier. Do not put the goods table here.
- line_fields: the columns of the goods / services table, in the printed order. Use [] if there is no table.
- Every item of fields and line_fields is {"label": ..., "type": ..., "hint": ..., "required": ...}; items of fields also have "zone".
    required: true when every document of this form always prints this value (for example invoice number, date, seller name, grand total; in the goods table the item name and the amount), otherwise false. Keep the number of required items small.
    label: the field name in natural Vietnamese, as a Vietnamese office worker would call it, short (at most 40 characters).
    type: one of text, number, money, date, percent, yesno. Use money for amounts of money, number for quantities and row numbers, date for dates, percent for tax rates.
    hint: a short Vietnamese phrase telling where to find it on this form, for example "Góc phải phía trên, sau chữ Số".
    zone: where the value is printed, one of:
        top_left   = header block on the left (seller name, address, tax code, phone)
        top_right  = header block on the right (form number, series, invoice number)
        subtitle   = lines right under the title (invoice date)
        info       = label / value lines between the title and the goods table (buyer, payment method, project)
        totals     = amounts printed under the goods table (subtotal, tax rate, tax, grand total)
        footer     = text lines after the totals (amount in words, bank transfer details)
        bottom     = small print at the very bottom (lookup code, digital signature, stamp text)
- notes: one or two Vietnamese sentences about layout details that could confuse reading (tilted scan, stamp over text, wrapped lines, punch holes), or "" if there are none.
Order the fields in reading order, top to bottom."""

def make_key(label, used):
    # ASCII key for the JSON answer, derived from the Vietnamese label
    text = unicodedata.normalize("NFKD", label.replace("đ", "d").replace("Đ", "D"))
    text = "".join(c for c in text if not unicodedata.combining(c))
    base = re.sub(r"[^a-z0-9]+", "_", text.lower()).strip("_")[:40] or "truong"
    key, n = base, 2
    while key in used or key == "line_items":
        key = f"{base}_{n}"
        n += 1
    used.add(key)
    return key

def clean_fields(raw, limit, with_zone=False):
    if not isinstance(raw, list):
        return [], None
    cleaned, used = [], set()
    for item in raw:
        if not isinstance(item, dict):
            continue
        label = str(item.get("label") or "").strip()
        if not label:
            continue  # blank rows in the form are ignored
        if len(label) > 60:
            return None, "Tên trường tối đa 60 ký tự."
        field_type = item.get("type") if item.get("type") in FIELD_TYPES else "text"
        hint = str(item.get("hint") or "").strip()[:200]
        key = item.get("key")
        if isinstance(key, str) and re.fullmatch(r"[a-z0-9_]{1,40}", key) and key not in used and key != "line_items":
            used.add(key)  # keep the key of an existing field so old results stay readable
        else:
            key = make_key(label, used)
        field = {"key": key, "label": label, "type": field_type, "hint": hint, "required": item.get("required") in (True, "true")}
        if with_zone:
            field["zone"] = item.get("zone") if item.get("zone") in ZONES else "info"
        cleaned.append(field)
    if len(cleaned) > limit:
        return None, f"Tối đa {limit} trường mỗi danh sách."
    return cleaned, None

def describe_field(field):
    line = f'    {field["key"]} (name on the form: "{field["label"]}"): {TYPE_INSTRUCTIONS[field["type"]]}.'
    if field.get("hint"):
        line += f' Where to find it: {field["hint"]}.'
    return line

def build_prompt(model):
    if model.get("prompt") and not model.get("fields") and not model.get("line_fields"):
        return model["prompt"]  # model saved before fields existed
    parts = [
        "You are reading one invoice-like document image (it may be Vietnamese, English or Arabic). Extract the data listed below.",
        "Respond with ONE JSON object inside a ```json code block and nothing else. Use exactly these keys.",
    ]
    fields = model.get("fields") or []
    if fields:
        parts.append("Fields printed once on the document:\n" + "\n".join(describe_field(f) for f in fields))
    line_fields = model.get("line_fields") or []
    if line_fields:
        parts.append(
            'Goods / services table: put every filled row of the table in an array named "line_items" (skip empty rows). '
            "Each row is an object with exactly these keys:\n" + "\n".join(describe_field(f) for f in line_fields)
        )
    else:
        parts.append("This document has no goods table, do not output line_items.")
    if model.get("notes"):
        parts.append("Extra notes from the user about this form:\n" + model["notes"])
    parts.append(RULES_TEXT)
    return "\n\n".join(parts)

### Review: which required fields could not be read, and why

# Sharpness below this (variance of the Laplacian on a 1000 px thumbnail) counts as blurry. Not calibrated yet, tune it with real scans.
BLUR_THRESHOLD = float(os.environ.get("BLUR_THRESHOLD", "150"))

def blur_score(image_bytes):
    # Higher = sharper. None when Pillow is missing or the image cannot be opened
    try:
        from PIL import Image, ImageFilter, ImageStat
        with Image.open(io.BytesIO(image_bytes)) as img:
            gray = img.convert("L")
            gray.thumbnail((1000, 1000))
            edges = gray.filter(ImageFilter.Kernel((3, 3), [0, 1, 0, 1, -4, 1, 0, 1, 0], scale=1, offset=128))
            return round(ImageStat.Stat(edges).var[0], 1)
    except Exception:
        return None

def is_empty(value):
    return value is None or (isinstance(value, str) and not value.strip())

def find_missing(schema, result):
    # Required fields (and goods table rows) that have no value
    missing = []
    for f in schema.get("fields") or []:
        if f.get("required") and is_empty(result.get(f["key"])):
            missing.append({"key": f["key"], "label": f["label"]})
    line_fields = schema.get("line_fields") or []
    if line_fields:
        rows = [r for r in (result.get("line_items") or []) if isinstance(r, dict)]
        if not rows:
            missing.append({"key": "line_items", "label": "Bảng hàng hóa"})
        for n, row in enumerate(rows):
            for f in line_fields:
                if f.get("required") and is_empty(row.get(f["key"])):
                    missing.append({"key": f"line_items.{n}.{f['key']}", "label": f"Dòng {n + 1}: {f['label']}"})
    return missing

def build_review(schema, result, model_quality, blur):
    missing = find_missing(schema, result)
    blurry = model_quality in ("blurry", "unreadable") and (blur is None or blur < BLUR_THRESHOLD)
    fields = schema.get("fields") or []
    empty_ratio = sum(is_empty(result.get(f["key"])) for f in fields) / len(fields) if fields else 0
    return {
        "state": "flagged" if missing else "ok",
        "reason": ("blur" if blurry else "other") if missing else None,
        "missing": missing,
        "model_quality": model_quality,
        "blur_score": blur,
        "wrong_model": len(fields) >= 4 and empty_ratio > 0.5,
    }

def analyze_result(schema, result, image_bytes):
    # Split the helper key out of the answer, then decide whether the invoice needs a human look
    if not isinstance(result, dict):
        raise ValueError("Kết quả của AI không phải một đối tượng JSON")
    result = dict(result)
    quality = result.pop("_image_quality", None)
    return result, build_review(schema, result, quality, blur_score(image_bytes))

def refresh_review(item, approve=False):
    # Recompute after a person edited the result. "approved" stays until the missing fields are filled in
    review = item.get("review") or {}
    missing = find_missing(item.get("schema") or {}, item["result"])
    review["missing"] = missing
    if not missing:
        review["state"], review["reason"] = "ok", None
    else:
        review["state"] = "approved" if approve or review.get("state") == "approved" else "flagged"
        review["reason"] = review.get("reason") or "other"
    item["review"] = review

def parse_vi_number(text):
    # "164.617.992" -> 164617992, "119,23" -> 119.23, "7.86" -> 7.86
    text = text.replace(" ", "").replace("%", "")
    if re.fullmatch(r"-?\d{1,3}(\.\d{3})+(,\d+)?", text) or "," in text:
        text = text.replace(".", "").replace(",", ".")
    value = float(text)
    return int(value) if value.is_integer() else value

def coerce_value(raw, field_type):
    # Turn what a person typed into the JSON value for this field type. Returns (value, error)
    if raw is None or (isinstance(raw, str) and not raw.strip()):
        return None, None
    if field_type == "yesno":
        text = str(raw).strip().lower()
        if text in ("có", "co", "true", "yes", "1"):
            return True, None
        if text in ("không", "khong", "false", "no", "0"):
            return False, None
        return None, "chọn Có hoặc Không"
    if field_type in ("number", "money", "percent"):
        if isinstance(raw, (int, float)) and not isinstance(raw, bool):
            return raw, None
        try:
            return parse_vi_number(str(raw).strip()), None
        except ValueError:
            return None, "phải là một số"
    text = str(raw).strip()[:1000]
    if field_type == "date":
        try:
            datetime.strptime(text, "%d/%m/%Y")
        except ValueError:
            return None, "ngày phải có dạng DD/MM/YYYY"
    return text, None

def manual_cells(item):
    # Keys of values that differ from what the AI read, so Excel can highlight what a person typed
    result, auto = item.get("result"), item.get("auto_result")
    if not isinstance(result, dict) or not isinstance(auto, dict):
        return set()
    same = lambda a, b: (None if is_empty(a) else a) == (None if is_empty(b) else b)
    schema = item.get("schema") or {}
    cells = {f["key"] for f in schema.get("fields") or [] if not same(result.get(f["key"]), auto.get(f["key"]))}
    rows, auto_rows = result.get("line_items") or [], auto.get("line_items") or []
    if len(rows) != len(auto_rows):
        cells.add("line_items")
    for n, row in enumerate(rows):
        base = auto_rows[n] if n < len(auto_rows) and isinstance(auto_rows[n], dict) else {}
        for f in schema.get("line_fields") or []:
            if isinstance(row, dict) and not same(row.get(f["key"]), base.get(f["key"])):
                cells.add(f"line_items.{n}.{f['key']}")
    return cells

def status_label(item):
    if item["status"] == "error":
        return "Lỗi"
    if item["status"] != "done":
        return "Đang xử lý"
    state = (item.get("review") or {}).get("state", "ok")
    if state == "flagged":
        return "Cần xem"
    if state == "approved":
        return "Duyệt để trống"
    if item.get("edited") and manual_cells(item):
        return "Nhập tay"
    return "Tự động"

MODELS_FILE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "models.json")
models_lock = threading.Lock()

def detect_mime_type(data):
    # Read the image type from its first bytes so Gemini gets the correct MIME type
    if data.startswith(b"\x89PNG"):
        return "image/png"
    if data.startswith(b"\xff\xd8\xff"):
        return "image/jpeg"
    if data[:4] == b"RIFF" and data[8:12] == b"WEBP":
        return "image/webp"
    return "image/png"

def allowed_file(mime_type):
    # Check if the MIME type is in the allowed list
    return mime_type in ALLOWED_MIME_TYPES

### Models: user-defined invoice types, each with its own extraction prompt, stored in models.json

def load_models():
    try:
        with open(MODELS_FILE, "r", encoding="utf-8") as f:
            return json.load(f)
    except FileNotFoundError:
        return []

def save_models(models):
    tmp_path = MODELS_FILE + ".tmp"
    with open(tmp_path, "w", encoding="utf-8") as f:
        json.dump(models, f, ensure_ascii=False, indent=2)
    os.replace(tmp_path, MODELS_FILE)

def find_model(key):
    # A model can be referenced by its id or by its name
    for model in load_models():
        if key in (model["id"], model["name"]):
            return model
    return None

def clean_signatures(raw):
    # Signature captions printed at the bottom of the form, accepted as a list or as one text with a caption per line
    if isinstance(raw, str):
        raw = raw.splitlines()
    if not isinstance(raw, list):
        return []
    return [str(s).strip()[:40] for s in raw if str(s).strip()][:4]

def read_model_payload(data):
    name = (data.get("name") or "").strip()
    if not name:
        return None, "Tên mô hình không được để trống."
    if len(name) > 60:
        return None, "Tên mô hình tối đa 60 ký tự."
    fields, error = clean_fields(data.get("fields"), 80, with_zone=True)
    if error:
        return None, error
    line_fields, error = clean_fields(data.get("line_fields"), 30)
    if error:
        return None, error
    if not fields and not line_fields:
        return None, "Hãy thêm ít nhất một trường cần đọc."
    notes = str(data.get("notes") or "").strip()[:2000]
    title = str(data.get("title") or "").strip()[:120]
    return {
        "name": name, "title": title, "signatures": clean_signatures(data.get("signatures")),
        "fields": fields, "line_fields": line_fields, "notes": notes,
    }, None

@app.route("/models", methods=["GET"])
def models_list():
    return jsonify({
        "models": load_models(),
        "types": FIELD_TYPES,
        "default_fields": DEFAULT_FIELDS,
        "default_line_fields": DEFAULT_LINE_FIELDS,
        "default_title": DEFAULT_TITLE,
        "default_signatures": DEFAULT_SIGNATURES,
    })

@app.route("/models/suggest", methods=["POST"])
def models_suggest():
    # Ask Gemini to propose the fields of a form from one sample image
    file = request.files.get("image")
    if file is None:
        return jsonify({"error": "Chưa chọn ảnh mẫu."}), 400
    image_bytes = file.read()
    try:
        model = genai.GenerativeModel(
            model_name="gemini-3.8-flash",
            generation_config=genai.types.GenerationConfig(temperature=0.2, response_mime_type="application/json"),
        )
        wait_for_rate_limit()
        response = model.generate_content([
            {"mime_type": detect_mime_type(image_bytes), "data": image_bytes}, "\n\n", SUGGEST_PROMPT,
        ])
        text = response.text.strip()
        if text.startswith("```"):
            text = text.split("```json")[-1].strip("` \n")
        data = json.loads(text)
    except Exception as e:
        return jsonify({"error": f"AI chưa gợi ý được: {e}"}), 502
    fields, _ = clean_fields(data.get("fields"), 80, with_zone=True)
    line_fields, _ = clean_fields(data.get("line_fields"), 30)
    return jsonify({
        "title": str(data.get("title") or "").strip()[:120],
        "signatures": clean_signatures(data.get("signatures")),
        "fields": fields or [],
        "line_fields": line_fields or [],
        "notes": str(data.get("notes") or "").strip()[:2000],
    })

@app.route("/models", methods=["POST"])
def models_create():
    payload, error = read_model_payload(request.get_json(silent=True) or {})
    if error:
        return jsonify({"error": error}), 400
    with models_lock:
        models = load_models()
        if any(m["name"].lower() == payload["name"].lower() for m in models):
            return jsonify({"error": "Đã có mô hình trùng tên."}), 409
        model = {"id": uuid.uuid4().hex[:8], **payload}
        models.append(model)
        save_models(models)
    return jsonify(model), 201

@app.route("/models/<model_id>", methods=["PUT"])
def models_update(model_id):
    payload, error = read_model_payload(request.get_json(silent=True) or {})
    if error:
        return jsonify({"error": error}), 400
    with models_lock:
        models = load_models()
        position = next((i for i, m in enumerate(models) if m["id"] == model_id), None)
        if position is None:
            return jsonify({"error": "Không tìm thấy mô hình."}), 404
        if any(m["id"] != model_id and m["name"].lower() == payload["name"].lower() for m in models):
            return jsonify({"error": "Đã có mô hình trùng tên."}), 409
        model = {"id": model_id, **payload}  # replaces the whole record, so an old raw prompt is dropped
        models[position] = model
        save_models(models)
    return jsonify(model)

@app.route("/models/<model_id>", methods=["DELETE"])
def models_delete(model_id):
    with models_lock:
        models = load_models()
        remaining = [m for m in models if m["id"] != model_id]
        if len(remaining) == len(models):
            return jsonify({"error": "Không tìm thấy mô hình."}), 404
        save_models(remaining)
    return jsonify({"ok": True})

@app.route("/", methods=["GET", "POST"])
def index():
    if request.method == "GET":
        return render_template("console.html")

    model = genai.GenerativeModel(
        model_name="gemini-3.8-flash",
        system_instruction=system_instruction,
        generation_config=genai.types.GenerationConfig(
            temperature=0.1,
        ),
    )

    try:
        if request.method == "POST":
            data = request.json  # Parse JSON payload
            base64_image = data.get("image")
            image_url = data.get("image_url")
            invoice_type = data.get("invoice_type")
            include_review = data.get("include_review") is True

            if not invoice_type:
                return jsonify({"error": "Missing 'invoice_type' parameter."}), 400

            model_def = find_model(invoice_type)
            if not model_def:
                return jsonify({"error": "Invalid invoice type. Please select a valid invoice type."}), 400

            unique_filename = f"uploaded_image_{uuid.uuid4().hex}.png"
            file_path = os.path.join(app.config["UPLOAD_FOLDER"], unique_filename)

            if base64_image:
                # Decode the base64 image and save it as a file
                try:
                    image_data = base64.b64decode(base64_image)
                    with open(file_path, "wb") as f:
                        f.write(image_data)
                except Exception as e:
                    return jsonify({"error": f"Error decoding base64 image: {str(e)}"}), 400
            elif image_url:
                # Download the image from the URL
                try:
                    response = requests.get(image_url, stream=True)
                    if response.status_code == 200:
                        with open(file_path, "wb") as f:
                            for chunk in response.iter_content(chunk_size=8192):
                                f.write(chunk)
                    else:
                        return jsonify({"error": f"Failed to download image from URL: HTTP {response.status_code}"}), 400
                except Exception as e:
                    return jsonify({"error": f"Error downloading image from URL: {str(e)}"}), 400
            else:
                return jsonify({"error": "Missing 'image' or 'image_url' parameter."}), 400

            # Generate content using the model
            prompt = build_prompt(model_def)
            with open(file_path, "rb") as f:
                image_bytes = f.read()
            file_to_upload = {"mime_type": detect_mime_type(image_bytes), "data": image_bytes}
            response = model.generate_content([file_to_upload, "\n\n", prompt])
            gemini_response = response.text
            if '```json' not in gemini_response:
                print("Unexpected response format:")
                print(gemini_response)
                return jsonify({"error": "Phản hồi của AI không chứa JSON."}), 502
            # Extract JSON part after "```json" and strip newlines
            json_part = gemini_response.split('```json')[-1].strip('` \n')
            parsed_json, review = analyze_result(model_def, json.loads(json_part), image_bytes)
            if include_review:
                parsed_json["_review"] = review  # opt-in so existing clients keep the same response shape
            print(json.dumps(parsed_json, indent=4, ensure_ascii=False))
            return parsed_json
    except Exception as e:
        return jsonify({"error": f"An error occurred: {str(e)}"}), 500

@app.route("/excel", methods=["POST"])
def export_excel():
    # Convert one extracted invoice (JSON) into an Excel file with two sheets
    invoice = request.json or {}
    line_items = invoice.pop("line_items", None) or []

    if request.args.get("layout") == "form":
        return send_workbook(build_form_workbook(invoice, line_items), invoice)

    wb = Workbook()
    header_sheet = wb.active
    header_sheet.title = "Hóa đơn"
    header_sheet.append(["Trường", "Giá trị"])
    for key, value in invoice.items():
        header_sheet.append([key, value if value is not None else ""])

    items_sheet = wb.create_sheet("Sản phẩm")
    columns = []
    for item in line_items:
        for key in item:
            if key not in columns:
                columns.append(key)
    items_sheet.append(columns)
    for item in line_items:
        items_sheet.append([item.get(col) if item.get(col) is not None else "" for col in columns])

    # Bold header rows and wrap long text
    for sheet in (header_sheet, items_sheet):
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        for col in sheet.columns:
            width = max(len(str(cell.value or "")) for cell in col)
            sheet.column_dimensions[col[0].column_letter].width = min(max(width + 2, 12), 60)

    return send_workbook(wb, invoice)

def build_form_workbook(invoice, line_items):
    # One sheet laid out like a printed invoice: seller block, buyer block, item table, total
    wb = Workbook()
    ws = wb.active
    ws.title = "Hóa đơn"

    bold = Font(bold=True)
    thin = Side(style="thin", color="999999")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill("solid", fgColor="E8EEF7")
    money = "#,##0.00"

    # Seller block (left) and invoice details (right)
    ws["A1"] = invoice.get("supplier_name") or ""
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = invoice.get("supplier_address") or ""
    ws["A3"] = f"MST: {invoice.get('supplier_vat') or ''}"

    details = [
        ("Số hóa đơn:", invoice.get("invoice_number")),
        ("Ngày:", invoice.get("invoice_date")),
        ("Khách hàng:", invoice.get("customer_name")),
        ("MST khách hàng:", invoice.get("customer_vat")),
    ]
    for row, (label, value) in enumerate(details, start=1):
        ws.cell(row=row, column=4, value=label).font = bold
        ws.cell(row=row, column=5, value=value or "")

    # Item table
    headers = ["HẠNG MỤC", "SỐ LƯỢNG", "ĐƠN GIÁ", "TỔNG CỘNG"]
    table_top = 6
    for col, title in enumerate(headers, start=1):
        cell = ws.cell(row=table_top, column=col, value=title)
        cell.font = bold
        cell.fill = header_fill
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    row = table_top + 1
    for item in line_items:
        name = item.get("item_description") or item.get("item_description_english") or item.get("item_description_arabic") or ""
        values = [name, item.get("quantity"), item.get("unit_price"), item.get("total_price")]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col, value="" if value is None else value)
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if col >= 2:
                cell.number_format = money
        row += 1

    # Total row
    ws.cell(row=row, column=3, value="Tổng cộng").font = bold
    total = ws.cell(row=row, column=4, value=invoice.get("total_amount_due"))
    total.font = bold
    total.number_format = money
    total.border = border

    widths = {"A": 42, "B": 12, "C": 16, "D": 18, "E": 30}
    for letter, width in widths.items():
        ws.column_dimensions[letter].width = width
    return wb

def send_workbook(wb, invoice):
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    # Keep only safe characters, invoice numbers can contain "/" or "#"
    safe_number = re.sub(r"[^\w-]", "_", str(invoice.get("invoice_number") or ""))
    filename = f"hoa_don_{safe_number or uuid.uuid4().hex[:8]}.xlsx"
    return send_file(
        buffer,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=filename,
    )

### Batch processing: many invoices, mixed types, processed in a background worker pool

BATCH_DIR = "batch_jobs/"
BATCH_WORKERS = int(os.environ.get("BATCH_WORKERS", "4"))  # parallel worker threads
BATCH_RPM = int(os.environ.get("BATCH_RPM", "5"))  # max Gemini requests per minute, 5 matches the free tier limit
MAX_RETRIES = 5
MAX_FILES_PER_JOB = 5000
MAX_MODELS_PER_JOB = 5

os.makedirs(BATCH_DIR, exist_ok=True)
executor = ThreadPoolExecutor(max_workers=BATCH_WORKERS)
jobs = {}  # job_id -> {"items": [...], "created": timestamp}
jobs_lock = threading.Lock()
_rate_lock = threading.Lock()
_next_request_at = [0.0]

# Jobs live in memory but are also written to batch_jobs/<job_id>/job.json so results and manual edits survive a restart
dirty_jobs = set()

def mark_dirty(job_id):
    dirty_jobs.add(job_id)

def flush_dirty():
    while dirty_jobs:
        try:
            job_id = dirty_jobs.pop()
        except KeyError:
            break
        with jobs_lock:
            job = jobs.get(job_id)
            text = json.dumps(job, ensure_ascii=False) if job else None
        job_dir = os.path.join(BATCH_DIR, job_id)
        if text is None or not os.path.isdir(job_dir):
            continue  # the job was deleted meanwhile
        path = os.path.join(job_dir, "job.json")
        try:
            with open(path + ".tmp", "w", encoding="utf-8") as f:
                f.write(text)
            os.replace(path + ".tmp", path)
        except OSError:
            dirty_jobs.add(job_id)  # try again on the next round

def saver_loop():
    while True:
        time.sleep(3)
        flush_dirty()

def load_jobs():
    for job_id in os.listdir(BATCH_DIR):
        try:
            with open(os.path.join(BATCH_DIR, job_id, "job.json"), "r", encoding="utf-8") as f:
                jobs[job_id] = json.load(f)
        except (OSError, ValueError):
            continue  # folders from before this feature have no job.json

load_jobs()
atexit.register(flush_dirty)
_background_started = [False]
_background_lock = threading.Lock()

@app.before_request
def start_background_once():
    # Done on the first request, not at import, so the debug reloader's parent process does not resume jobs a second time
    if _background_started[0]:
        return
    with _background_lock:
        if _background_started[0]:
            return
        _background_started[0] = True
        threading.Thread(target=saver_loop, daemon=True).start()
        for job in list(jobs.values()):
            for item in job["items"]:
                if item["status"] == "pending":
                    executor.submit(process_item, item)  # interrupted by a restart, read it again

def wait_for_rate_limit():
    # Reserve the next time slot so all worker threads together stay under BATCH_RPM
    interval = 60.0 / BATCH_RPM
    with _rate_lock:
        now = time.time()
        wait = max(0.0, _next_request_at[0] - now)
        _next_request_at[0] = max(now, _next_request_at[0]) + interval
    if wait > 0:
        time.sleep(wait)

def extract_invoice(image_bytes, prompt):
    # Same extraction as the single-invoice endpoint, but raises on failure instead of hiding it
    model = genai.GenerativeModel(
        model_name="gemini-3.8-flash",
        system_instruction=system_instruction,
        generation_config=genai.types.GenerationConfig(temperature=0.1),
    )
    file_part = {"mime_type": detect_mime_type(image_bytes), "data": image_bytes}
    wait_for_rate_limit()
    response = model.generate_content([file_part, "\n\n", prompt])
    text = response.text
    if "```json" not in text:
        raise ValueError("Phản hồi của AI không chứa JSON")
    json_part = text.split("```json")[-1].strip("` \n")
    return json.loads(json_part)

def extract_with_retry(image_bytes, prompt):
    last_error = None
    for attempt in range(MAX_RETRIES):
        try:
            return extract_invoice(image_bytes, prompt)
        except Exception as e:
            last_error = e
            if "429" in str(e) or "quota" in str(e).lower():
                time.sleep(65)  # quota error: wait out the one-minute window
            else:
                time.sleep(2 ** (attempt + 1))
    raise last_error

def process_item(item):
    if item.get("deleted"):
        return
    try:
        with open(item["path"], "rb") as f:
            image_bytes = f.read()
    except OSError:
        return  # the file was deleted while the item was waiting in the queue
    try:
        result = extract_with_retry(image_bytes, item["prompt"])
        result, review = analyze_result(item.get("schema") or {}, result, image_bytes)
        with jobs_lock:
            if not item.get("deleted"):
                item["result"] = result
                item["auto_result"] = copy.deepcopy(result)  # what the AI read, to tell later edits apart
                item["review"] = review
                item["status"] = "done"
    except Exception as e:
        with jobs_lock:
            if not item.get("deleted"):
                item["error"] = str(e)
                item["status"] = "error"
    mark_dirty(item["job_id"])

def find_item(job, index):
    # Items keep their original index even after others are deleted, so look up by value
    return next((i for i in job["items"] if i["index"] == index), None)

def remove_files(items):
    for item in items:
        item["deleted"] = True
        try:
            os.remove(item["path"])
        except OSError:
            pass

def job_summary(job):
    items = job["items"]
    done = sum(1 for i in items if i["status"] == "done")
    errors = sum(1 for i in items if i["status"] == "error")
    by_type = {}
    for i in items:
        entry = by_type.setdefault(i["invoice_type"], {"total": 0, "done": 0, "errors": 0})
        entry["total"] += 1
        if i["status"] == "done":
            entry["done"] += 1
        elif i["status"] == "error":
            entry["errors"] += 1
    review = {"flagged": 0, "flagged_blur": 0, "flagged_other": 0, "approved": 0, "manual": 0}
    for i in items:
        if i["status"] != "done":
            continue
        info = i.get("review") or {}
        if info.get("state") == "flagged":
            review["flagged"] += 1
            review["flagged_blur" if info.get("reason") == "blur" else "flagged_other"] += 1
        elif info.get("state") == "approved":
            review["approved"] += 1
        elif i.get("edited") and manual_cells(i):
            review["manual"] += 1
    return {
        "total": len(items),
        "done": done,
        "errors": errors,
        "finished": done + errors == len(items),
        "by_type": by_type,
        "review": review,
    }

@app.route("/batch", methods=["GET"])
def batch_page():
    return redirect(url_for("index"))

@app.route("/batch/<job_id>/results", methods=["GET"])
def batch_results(job_id):
    # Lightweight rows for the results table, full data of one invoice comes from /item/<index>
    with jobs_lock:
        job = jobs.get(job_id)
        if job is None:
            return jsonify({"error": "Không tìm thấy job."}), 404
        rows = []
        for i in job["items"]:
            result = i["result"] or {}
            summary = []
            for field in (i.get("schema") or {}).get("fields", []):
                value = result.get(field["key"])
                if value is None or value == "" or isinstance(value, (dict, list)):
                    continue
                summary.append({"label": field["label"], "value": value, "type": field["type"]})
                if len(summary) == 4:
                    break
            rows.append({
                "index": i["index"],
                "filename": i["filename"],
                "invoice_type": i["invoice_type"],
                "status": i["status"],
                "status_label": status_label(i),
                "review_state": (i.get("review") or {}).get("state"),
                "summary": summary,
                "line_count": len(result.get("line_items") or []),
                "error": i["error"],
            })
    return jsonify(rows)

def item_detail(item):
    # Call with jobs_lock held
    return {
        "filename": item["filename"],
        "invoice_type": item["invoice_type"],
        "status": item["status"],
        "status_label": status_label(item),
        "result": item["result"],
        "error": item["error"],
        "schema": item.get("schema") or {"fields": [], "line_fields": []},
        "review": item.get("review"),
        "manual": sorted(manual_cells(item)),
    }

@app.route("/batch/<job_id>/item/<int:index>", methods=["GET"])
def batch_item(job_id, index):
    with jobs_lock:
        job = jobs.get(job_id)
        item = find_item(job, index) if job else None
        if item is None:
            return jsonify({"error": "Không tìm thấy hóa đơn."}), 404
        return jsonify(item_detail(item))

@app.route("/batch/<job_id>/item/<int:index>", methods=["PATCH"])
def batch_item_update(job_id, index):
    # A person fills in or corrects values, optionally approving the invoice even if some required fields stay empty
    payload = request.get_json(silent=True) or {}
    with jobs_lock:
        job = jobs.get(job_id)
        item = find_item(job, index) if job else None
        if item is None:
            return jsonify({"error": "Không tìm thấy hóa đơn."}), 404
        if item["status"] != "done" or not isinstance(item.get("result"), dict):
            return jsonify({"error": "Hóa đơn này chưa có kết quả để sửa."}), 409
        schema = item.get("schema") or {}
        result = copy.deepcopy(item["result"])
        errors = []

        typed = payload.get("fields")
        for f in schema.get("fields") or []:
            if isinstance(typed, dict) and f["key"] in typed:
                value, error = coerce_value(typed[f["key"]], f["type"])
                if error:
                    errors.append(f"{f['label']}: {error}")
                else:
                    result[f["key"]] = value

        line_fields = schema.get("line_fields") or []
        if line_fields and isinstance(payload.get("line_items"), list):
            old_rows = result.get("line_items") or []
            new_rows = []
            for n, raw_row in enumerate(payload["line_items"]):
                if not isinstance(raw_row, dict):
                    continue
                row = dict(old_rows[n]) if n < len(old_rows) and isinstance(old_rows[n], dict) else {}
                for f in line_fields:
                    if f["key"] in raw_row:
                        value, error = coerce_value(raw_row[f["key"]], f["type"])
                        if error:
                            errors.append(f"Dòng {n + 1}, {f['label']}: {error}")
                        else:
                            row[f["key"]] = value
                if any(not is_empty(row.get(f["key"])) for f in line_fields):
                    new_rows.append(row)  # a row left completely blank is dropped
            result["line_items"] = new_rows

        if errors:
            return jsonify({"error": "\n".join(errors[:5])}), 400
        item["result"] = result
        item["edited"] = True
        refresh_review(item, approve=payload.get("approve") is True)
        mark_dirty(job_id)
        return jsonify(item_detail(item))

@app.route("/batch/<job_id>/item/<int:index>", methods=["DELETE"])
def batch_item_delete(job_id, index):
    with jobs_lock:
        job = jobs.get(job_id)
        item = find_item(job, index) if job else None
        if item is None:
            return jsonify({"error": "Không tìm thấy hóa đơn."}), 404
        job["items"].remove(item)
        remove_files([item])
    mark_dirty(job_id)
    return jsonify({"ok": True})

@app.route("/batch/<job_id>", methods=["DELETE"])
def batch_delete(job_id):
    # With ?model=<name> delete only that model's results, otherwise delete the whole job
    model_name = request.args.get("model")
    with jobs_lock:
        job = jobs.get(job_id)
        if job is None:
            return jsonify({"error": "Không tìm thấy job."}), 404
        if model_name:
            targets = [i for i in job["items"] if i["invoice_type"] == model_name]
            job["items"] = [i for i in job["items"] if i["invoice_type"] != model_name]
        else:
            targets = list(job["items"])
            job["items"] = []
            del jobs[job_id]
        remove_files(targets)
    if not model_name:
        dirty_jobs.discard(job_id)
        shutil.rmtree(os.path.join(BATCH_DIR, job_id), ignore_errors=True)
    else:
        mark_dirty(job_id)
    return jsonify({"ok": True, "deleted": len(targets)})

@app.route("/batch/<job_id>/image/<int:index>", methods=["GET"])
def batch_image(job_id, index):
    # Serve the uploaded scan so the detail panel can show it next to the extracted data
    with jobs_lock:
        job = jobs.get(job_id)
        item = find_item(job, index) if job else None
        if item is None:
            return jsonify({"error": "Không tìm thấy ảnh."}), 404
        path = item["path"]
    with open(path, "rb") as f:
        data = f.read()
    return send_file(io.BytesIO(data), mimetype=detect_mime_type(data))

@app.route("/batch/start", methods=["POST"])
def batch_start():
    files = request.files.getlist("files")
    types = request.form.getlist("types")  # one model id per file, same order as files
    default_type = request.form.get("invoice_type")
    models_by_id = {m["id"]: m for m in load_models()}

    if not files:
        return jsonify({"error": "Chưa chọn file nào."}), 400
    if len(files) > MAX_FILES_PER_JOB:
        return jsonify({"error": f"Tối đa {MAX_FILES_PER_JOB} file mỗi lần."}), 400

    chosen_types = set(types) if types else {default_type}
    if len(chosen_types) > MAX_MODELS_PER_JOB:
        return jsonify({"error": f"Chỉ được chọn tối đa {MAX_MODELS_PER_JOB} mô hình mỗi lần."}), 400

    for model_id in chosen_types:
        if model_id not in models_by_id:
            return jsonify({"error": "Mô hình không tồn tại hoặc đã bị xóa. Hãy tải lại trang."}), 400

    job_id = uuid.uuid4().hex
    job_dir = os.path.join(BATCH_DIR, job_id)
    os.makedirs(job_dir, exist_ok=True)

    items = []
    for index, file in enumerate(files):
        model = models_by_id[types[index] if index < len(types) else default_type]
        path = os.path.join(job_dir, f"{index:05d}.bin")
        file.save(path)
        items.append({
            "job_id": job_id,
            "index": index,
            "filename": file.filename,
            "invoice_type": model["name"],
            # Copied now so editing or deleting the model later does not affect this job
            "prompt": build_prompt(model),
            "schema": {
                "fields": model.get("fields") or [], "line_fields": model.get("line_fields") or [],
                "title": model.get("title") or "", "signatures": model.get("signatures") or [],
            },
            "path": path,
            "status": "pending",
            "result": None,
            "error": None,
        })

    with jobs_lock:
        jobs[job_id] = {"items": items, "created": time.time()}
    mark_dirty(job_id)
    for item in items:
        executor.submit(process_item, item)

    return jsonify({"job_id": job_id, "total": len(items)})

@app.route("/batch/<job_id>/status", methods=["GET"])
def batch_status(job_id):
    with jobs_lock:
        job = jobs.get(job_id)
        if job is None:
            return jsonify({"error": "Không tìm thấy job."}), 404
        summary = job_summary(job)
        failed = [
            {"filename": i["filename"], "error": i["error"]}
            for i in job["items"] if i["status"] == "error"
        ]
    summary["failed"] = failed[:200]  # only the first failures, the rest are in the Excel file
    return jsonify(summary)

@app.route("/batch/<job_id>/excel", methods=["GET"])
def batch_excel(job_id):
    with jobs_lock:
        job = jobs.get(job_id)
        if job is None:
            return jsonify({"error": "Không tìm thấy job."}), 404
        if not job_summary(job)["finished"]:
            return jsonify({"error": "Job chưa xử lý xong."}), 409
        items = [dict(i) for i in job["items"]]  # copy so the workbook is built without holding the lock

    if request.args.get("format") == "table":
        return send_workbook(build_batch_workbook(items), {"invoice_number": f"bang_{job_id[:8]}"})
    return send_workbook(build_layout_workbook(items), {"invoice_number": f"batch_{job_id[:8]}"})

NUMBER_FORMATS = {"money": "#,##0", "percent": '0.##"%"', "date": "dd/mm/yyyy"}
MANUAL_FILL = PatternFill("solid", fgColor="FFF2B3")  # values a person typed or corrected

def excel_value(value, field_type):
    # Turn a JSON value into something Excel stores with the right type
    if value is None or value == "":
        return None
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    if field_type in ("number", "money", "percent") and not isinstance(value, (int, float)):
        try:
            return float(str(value).replace(" ", ""))
        except ValueError:
            return value
    if field_type == "date":
        try:
            return datetime.strptime(str(value), "%d/%m/%Y")
        except ValueError:
            return value
    if field_type == "yesno":
        return "Có" if value is True else "Không" if value is False else value
    return value

def append_typed(sheet, cells):
    # cells: list of (value, field_type or None)
    sheet.append([excel_value(value, field_type) for value, field_type in cells])
    for column, (_, field_type) in enumerate(cells, start=1):
        number_format = NUMBER_FORMATS.get(field_type)
        if number_format:
            sheet.cell(row=sheet.max_row, column=column).number_format = number_format

def style_sheet(sheet):
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E8EEF7")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in range(1, sheet.max_column + 1):
        # Width from the longest value in the first 500 rows, enough for a quick look
        longest = 0
        for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 500), min_col=column, max_col=column):
            value = row[0].value
            if value is not None:
                longest = max(longest, len(str(value)))
        sheet.column_dimensions[sheet.cell(row=1, column=column).column_letter].width = min(max(longest + 2, 10), 50)

def sheet_title(name, used):
    # Excel sheet names: at most 31 characters, none of []:*?/\, unique
    base = re.sub(r"[\[\]:*?/\\]", "_", name).strip()[:31] or "Mô hình"
    title, n = base, 2
    while title.lower() in used:
        suffix = f" ({n})"
        title = base[:31 - len(suffix)] + suffix
        n += 1
    used.add(title.lower())
    return title

def build_batch_workbook(items):
    wb = Workbook()
    summary_sheet = wb.active
    summary_sheet.title = "Tổng hợp"
    used_titles = {"tổng hợp", "chi tiết hàng"}
    summary_sheet.append(["STT", "Tên file", "Mô hình", "Trạng thái", "Xử lý", "Lỗi"])

    groups = {}  # model name -> [(position, item)], keeps the order of the files
    for position, item in enumerate(items, start=1):
        status = "Thành công" if item["status"] == "done" else "Lỗi"
        summary_sheet.append([position, item["filename"], item["invoice_type"], status, status_label(item), item["error"] or ""])
        groups.setdefault(item["invoice_type"], []).append((position, item))
    style_sheet(summary_sheet)

    # One sheet per model: the columns are the fields the user defined, in the same order
    detail_labels = []
    for model_name, entries in groups.items():
        schema = entries[0][1].get("schema") or {"fields": [], "line_fields": []}
        fields = schema["fields"]
        sheet = wb.create_sheet(sheet_title(model_name, used_titles))
        sheet.append(["STT", "Tên file", "Trạng thái", "Xử lý"] + [f["label"] for f in fields] + ["Lỗi"])
        for position, item in entries:
            result = item["result"] or {}
            status = "Thành công" if item["status"] == "done" else "Lỗi"
            cells = [(position, None), (item["filename"], None), (status, None), (status_label(item), None)]
            cells += [(result.get(f["key"]), f["type"]) for f in fields]
            cells.append((item["error"] or "", None))
            append_typed(sheet, cells)
            manual = manual_cells(item)
            for column, f in enumerate(fields, start=5):
                if f["key"] in manual:
                    sheet.cell(row=sheet.max_row, column=column).fill = MANUAL_FILL
        style_sheet(sheet)
        for field in schema["line_fields"]:
            if field["label"] not in detail_labels:
                detail_labels.append(field["label"])

    # Goods table of every invoice in one sheet, columns matched by the field name
    if detail_labels:
        detail_sheet = wb.create_sheet("Chi tiết hàng")
        detail_sheet.append(["STT", "Tên file", "Mô hình", "Dòng"] + detail_labels)
        for position, item in enumerate(items, start=1):
            line_fields = {f["label"]: f for f in (item.get("schema") or {}).get("line_fields", [])}
            manual = manual_cells(item)
            for row_number, line in enumerate((item["result"] or {}).get("line_items") or [], start=1):
                cells = [(position, None), (item["filename"], None), (item["invoice_type"], None), (row_number, None)]
                for label in detail_labels:
                    field = line_fields.get(label)
                    cells.append((line.get(field["key"]), field["type"]) if field else (None, None))
                append_typed(detail_sheet, cells)
                for column, label in enumerate(detail_labels, start=5):
                    field = line_fields.get(label)
                    if field and f"line_items.{row_number - 1}.{field['key']}" in manual:
                        detail_sheet.cell(row=detail_sheet.max_row, column=column).fill = MANUAL_FILL
        style_sheet(detail_sheet)
    return wb

### Excel that reproduces the layout of the original form, everything on one sheet

LAYOUT_SIDE = Side(style="thin", color="888888")
LAYOUT_BOX = Border(left=LAYOUT_SIDE, right=LAYOUT_SIDE, top=LAYOUT_SIDE, bottom=LAYOUT_SIDE)
LAYOUT_HEAD_FILL = PatternFill("solid", fgColor="DCE6F5")
LAYOUT_BLUE = "1F4E8C"
LAYOUT_RED = "B42318"
LAYOUT_MUTED = "666666"
STT_LABELS = {"STT", "#", "TT"}

def vi_number(value):
    # 1234567.5 -> "1.234.567,5" as Vietnamese forms print numbers
    if float(value).is_integer():
        return f"{int(value):,}".replace(",", ".")
    text = f"{value:,.2f}".rstrip("0").rstrip(".")
    return text.replace(",", "_").replace(".", ",").replace("_", ".")

def vi_text(value, field_type):
    if value is None or value == "":
        return ""
    if isinstance(value, bool) or field_type == "yesno":
        return "Có" if value is True else "Không" if value is False else str(value)
    if isinstance(value, (int, float)) and field_type in ("number", "money", "percent"):
        return vi_number(value) + ("%" if field_type == "percent" else "")
    return str(value)

def vi_long_date(value):
    # "15/06/2026" -> "Ngày 15 tháng 06 năm 2026", the way the date line is printed under the title
    try:
        parsed = datetime.strptime(str(value), "%d/%m/%Y")
        return f"Ngày {parsed.day:02d} tháng {parsed.month:02d} năm {parsed.year}"
    except ValueError:
        return str(value or "")

def rich_line(label, value_text, italic=False):
    if not value_text:
        return f"{label}:"
    return CellRichText(
        TextBlock(InlineFont(color=LAYOUT_MUTED), f"{label}: "),
        TextBlock(InlineFont(b=not italic, i=italic), value_text),
    )

def put(ws, row, first_col, last_col, value, font=None, align=None, fill=None, border=None, number_format=None):
    cell = ws.cell(row=row, column=first_col, value=value)
    if font:
        cell.font = font
    if align:
        cell.alignment = align
    if fill:
        cell.fill = fill
    if number_format:
        cell.number_format = number_format
    if last_col > first_col:
        ws.merge_cells(start_row=row, start_column=first_col, end_row=row, end_column=last_col)
    if border:
        for column in range(first_col, last_col + 1):
            ws.cell(row=row, column=column).border = border
    return cell

def render_invoice(ws, row, item, ctx):
    # Draw one invoice starting at `row`, return the next free row
    schema = item.get("schema") or {}
    fields = schema.get("fields") or []
    line_fields = schema.get("line_fields") or []
    k = max(len(line_fields), 4)  # number of columns this form uses
    result = item.get("result")
    wrap_top = Alignment(wrap_text=True, vertical="top")
    manual = manual_cells(item)
    fill_of = lambda key: MANUAL_FILL if key in manual else None

    def note_height(at_row, text_len, first_col, last_col, base=16):
        ctx["heights"].append((at_row, text_len, first_col, last_col, base))

    if item["status"] != "done" or not result:
        message = f"{item['filename']}: không đọc được ({item.get('error') or 'chưa xử lý xong'})"
        put(ws, row, 1, k, message, font=Font(color=LAYOUT_RED), align=wrap_top)
        note_height(row, len(message), 1, k)
        return row + 2

    zones = {zone: [] for zone in ZONES}
    for field in fields:
        zone = field.get("zone") if field.get("zone") in zones else "info"
        zones[zone].append((field, result.get(field["key"])))

    # Header: seller block on the left, form number block on the right
    left, right = zones["top_left"], zones["top_right"]
    split = max(2, k - 2)
    for i, (field, value) in enumerate(left):
        text = vi_text(value, field["type"])
        if i == 0:
            put(ws, row + i, 1, split, text, font=Font(bold=True, size=13, color=LAYOUT_BLUE), align=wrap_top, fill=fill_of(field["key"]))
            note_height(row + i, len(text) * 1.2, 1, split, 20)
        else:
            put(ws, row + i, 1, split, rich_line(field["label"], text), align=wrap_top, fill=fill_of(field["key"]))
            note_height(row + i, len(field["label"]) + len(text) + 2, 1, split)
    for i, (field, value) in enumerate(right):
        text = vi_text(value, field["type"])
        put(ws, row + i, split + 1, k, rich_line(field["label"], text), align=Alignment(horizontal="right", vertical="top", wrap_text=True), fill=fill_of(field["key"]))
        note_height(row + i, len(field["label"]) + len(text) + 2, split + 1, k)
    row += max(len(left), len(right)) + (1 if left or right else 0)

    # Title and the lines right under it
    title = schema.get("title") or item["invoice_type"]
    put(ws, row, 1, k, title, font=Font(bold=True, size=16, color=LAYOUT_RED), align=Alignment(horizontal="center", vertical="center", wrap_text=True))
    ws.row_dimensions[row].height = 30
    row += 1
    for field, value in zones["subtitle"]:
        shown = vi_long_date(value) if field["type"] == "date" else f"{field['label']}: {vi_text(value, field['type'])}"
        put(ws, row, 1, k, shown, font=Font(italic=True), align=Alignment(horizontal="center"), fill=fill_of(field["key"]))
        row += 1
    row += 1

    # Label / value lines between the title and the table (buyer, payment method...)
    for field, value in zones["info"]:
        text = vi_text(value, field["type"])
        put(ws, row, 1, k, rich_line(field["label"], text), align=wrap_top, fill=fill_of(field["key"]))
        note_height(row, len(field["label"]) + len(text) + 2, 1, k)
        row += 1
    if zones["info"]:
        row += 1

    # Goods table
    if line_fields:
        for column, field in enumerate(line_fields, start=1):
            put(ws, row, column, column, field["label"], font=Font(bold=True), fill=LAYOUT_HEAD_FILL, border=LAYOUT_BOX,
                align=Alignment(horizontal="center", vertical="center", wrap_text=True))
            ctx["widths"][column] = max(ctx["widths"].get(column, 8), min(len(field["label"]) + 2, 18))
        ws.row_dimensions[row].height = 22
        row += 1
        for line_number, line in enumerate(result.get("line_items") or []):
            for column, field in enumerate(line_fields, start=1):
                raw = line.get(field["key"])
                numeric = field["type"] in ("number", "money", "percent")
                if field["label"].strip().upper() in STT_LABELS:
                    alignment = Alignment(horizontal="center", vertical="top")
                elif numeric:
                    alignment = Alignment(horizontal="right", vertical="top")
                else:
                    alignment = wrap_top
                put(ws, row, column, column, excel_value(raw, field["type"]), border=LAYOUT_BOX, align=alignment,
                    number_format=NUMBER_FORMATS.get(field["type"]), fill=fill_of(f"line_items.{line_number}.{field['key']}"))
                ctx["widths"][column] = max(ctx["widths"].get(column, 8), min(len(vi_text(raw, field["type"])) + 2, 48))
            row += 1

    # Totals, right aligned under the last columns of the table
    totals = zones["totals"]
    label_first = max(1, k - 3)
    for i, (field, value) in enumerate(totals):
        last = i == len(totals) - 1
        put(ws, row, label_first, k - 1, field["label"] + ":", font=Font(bold=last), align=Alignment(horizontal="right"))
        put(ws, row, k, k, excel_value(value, field["type"]), font=Font(bold=last), border=LAYOUT_BOX,
            fill=fill_of(field["key"]) or (LAYOUT_HEAD_FILL if last else None), align=Alignment(horizontal="right"),
            number_format=NUMBER_FORMATS.get(field["type"]))
        row += 1
    if totals:
        row += 1

    # Text lines after the totals (amount in words...)
    for field, value in zones["footer"]:
        text = vi_text(value, field["type"])
        put(ws, row, 1, k, rich_line(field["label"], text, italic=True), align=wrap_top, fill=fill_of(field["key"]))
        note_height(row, len(field["label"]) + len(text) + 2, 1, k)
        row += 1
    if zones["footer"]:
        row += 1

    # Signature captions spread evenly across the table width
    signatures = (schema.get("signatures") or [])[:k]
    if signatures:
        count = len(signatures)
        for index, caption in enumerate(signatures):
            first, last = index * k // count + 1, (index + 1) * k // count
            put(ws, row, first, last, caption, font=Font(bold=True), align=Alignment(horizontal="center"))
            put(ws, row + 1, first, last, "(Ký, ghi rõ họ tên)", font=Font(italic=True, color=LAYOUT_MUTED, size=9), align=Alignment(horizontal="center"))
        ws.row_dimensions[row + 2].height = 50
        row += 3

    # Small print at the very bottom (lookup code, digital signature, stamp)
    for field, value in zones["bottom"]:
        text = vi_text(value, field["type"])
        put(ws, row, 1, k, f"{field['label']}: {text}", font=Font(size=9, color=LAYOUT_MUTED), align=Alignment(horizontal="center", wrap_text=True), fill=fill_of(field["key"]))
        note_height(row, len(field["label"]) + len(text) + 2, 1, k, 14)
        row += 1
    return row + 1

def build_layout_workbook(items):
    wb = Workbook()
    ws = wb.active
    ws.title = "Hóa đơn"
    ws.sheet_view.showGridLines = False
    ctx = {"widths": {}, "heights": []}

    row = 1
    total = len(items)
    for position, item in enumerate(items, start=1):
        label = status_label(item)
        if total > 1 or label != "Tự động":
            banner = f"Hóa đơn {position}/{total} · {item['filename']} · {item['invoice_type']}" + ("" if label == "Tự động" else f" · {label}")
            put(ws, row, 1, 4, banner, font=Font(size=9, color=LAYOUT_MUTED), fill=PatternFill("solid", fgColor="EEF1F6"))
            row += 2
        row = render_invoice(ws, row, item, ctx)
        if total > 1 and position < total:
            ws.row_breaks.append(Break(id=row - 1))  # one invoice per printed page
            row += 1

    for column in range(1, max(ctx["widths"], default=4) + 1):
        ws.column_dimensions[get_column_letter(column)].width = ctx["widths"].get(column, 12)
    # Excel does not grow merged cells by itself, so set their height from the final column widths
    for at_row, text_len, first_col, last_col, base in ctx["heights"]:
        span = sum(ctx["widths"].get(column, 12) for column in range(first_col, last_col + 1))
        lines = max(1, math.ceil(text_len / max(span * 1.05, 10)))
        ws.row_dimensions[at_row].height = max(base, 15 * lines)

    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    return wb

@app.route("/batch/<job_id>/item/<int:index>/excel", methods=["GET"])
def batch_item_excel(job_id, index):
    with jobs_lock:
        job = jobs.get(job_id)
        item = find_item(job, index) if job else None
        if item is None:
            return jsonify({"error": "Không tìm thấy hóa đơn."}), 404
        item = dict(item)
    stem = os.path.splitext(item["filename"] or "")[0]
    return send_workbook(build_layout_workbook([item]), {"invoice_number": stem})

if __name__ == "__main__":
    app.run(debug=True)
